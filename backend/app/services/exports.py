"""Excel / CSV exports: dispatch reports, courier hand-over manifests, OMS dispatch-update file."""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import OmsOrder, Scan
from ..services.scanning import scan_order
from ..timeutil import to_local

SCAN_COLUMNS = [
    ("S.No", 6), ("Dispatch Date", 12), ("Scanned At", 18), ("Sales Channel", 28), ("Marketplace", 16),
    ("AWB / Tracking", 22), ("Courier", 16), ("Order ID", 22), ("Sub Order IDs", 26), ("Invoice No", 20),
    ("SKUs", 30), ("Qty", 6), ("Amount", 10), ("Payment", 10), ("Buyer", 20), ("City", 14), ("Pincode", 9),
    ("Result", 11), ("Flags", 24), ("Alert", 30), ("Scanned By", 16), ("Station", 12), ("Manifest", 10),
    ("OMS Update", 11),
]


def _skus(o: OmsOrder | None) -> str:
    if not o:
        return ""
    import json

    try:
        items = json.loads(o.items_json or "[]")
    except ValueError:
        return ""
    return ", ".join(f"{i.get('sku')} x{i.get('qty')}" for i in items)


def _od(s: Scan) -> dict:
    """Order details of a scan (works after the order has been pruned from the working set)."""
    return scan_order(s) or {}


def _od_skus(od: dict) -> str:
    return ", ".join(f"{i.get('sku')} x{i.get('qty')}" for i in od.get("items") or [])


def scan_row(i: int, s: Scan) -> list:
    od = _od(s)
    return [
        i, s.dispatch_date.strftime("%d-%m-%Y"), to_local(s.scanned_at).strftime("%d-%m-%Y %H:%M:%S"),
        s.channel.name if s.channel else "", s.channel.marketplace if s.channel else "",
        s.tracking_raw, od.get("courier", ""), od.get("channel_order_id", ""),
        ", ".join(od.get("sub_order_ids") or []), od.get("invoice_id", ""), _od_skus(od), od.get("total_qty", ""),
        od.get("total_amount", ""), od.get("order_type", ""), od.get("buyer_name", ""),
        od.get("buyer_city", ""), od.get("buyer_pincode", ""), s.result, s.flags.replace(",", ", "),
        s.alert, (s.user.full_name or s.user.username) if s.user else "", s.station,
        s.manifest_id or "", s.oms_update_status,
    ]


def _style_sheet(ws, widths: Sequence[int], header_row: int) -> None:
    head_fill = PatternFill("solid", fgColor="1F2937")
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{max(ws.max_row, header_row)}"


def scans_xlsx(scans: Iterable[Scan], title: str, summary: list[tuple[str, object]] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dispatch"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    r = 2
    for k, v in summary or []:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        r += 1
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c for c, _ in SCAN_COLUMNS])
    warn_fill = PatternFill("solid", fgColor="FEF3C7")
    alert_fill = PatternFill("solid", fgColor="FEE2E2")
    for i, s in enumerate(scans, start=1):
        ws.append(scan_row(i, s))
        if s.alert:
            for c in ws[ws.max_row]:
                c.fill = alert_fill
        elif s.result != "OK":
            for c in ws[ws.max_row]:
                c.fill = warn_fill
    _style_sheet(ws, [w for _, w in SCAN_COLUMNS], header_row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def manifest_xlsx(scans: Sequence[Scan], *, channel_name: str, dispatch_date: str, manifest_no: str,
                  closed_at: datetime | None) -> bytes:
    """Courier hand-over sheet with signature block."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Manifest"
    ws.append([f"Dispatch Manifest - {channel_name}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Manifest No", manifest_no])
    ws.append(["Dispatch Date", dispatch_date])
    ws.append(["Total Shipments", len(scans)])
    ws.append(["Generated", to_local(closed_at).strftime("%d-%m-%Y %H:%M") if closed_at else "OPEN (not closed)"])
    couriers: dict[str, int] = {}
    for s in scans:
        k = _od(s).get("courier") or "Unknown"
        couriers[k] = couriers.get(k, 0) + 1
    ws.append(["By Courier", ", ".join(f"{k}: {v}" for k, v in sorted(couriers.items()))])
    for row in range(2, 7):
        ws.cell(row=row, column=1).font = Font(bold=True)
    ws.append([])
    cols = [("S.No", 6), ("AWB / Tracking", 24), ("Courier", 18), ("Order ID", 24), ("Invoice No", 20),
            ("Qty", 6), ("Payment", 10), ("Pincode", 10), ("Scanned At", 18)]
    header_row = ws.max_row + 1
    ws.append([c for c, _ in cols])
    for i, s in enumerate(scans, start=1):
        od = _od(s)
        ws.append([i, s.tracking_raw, od.get("courier", ""), od.get("channel_order_id", ""),
                   od.get("invoice_id", ""), od.get("total_qty", ""), od.get("order_type", ""),
                   od.get("buyer_pincode", ""), to_local(s.scanned_at).strftime("%d-%m-%Y %H:%M")])
    _style_sheet(ws, [w for _, w in cols], header_row)
    ws.append([])
    ws.append(["Handed over by (name / sign):", "", "", "Received by courier (name / sign):"])
    ws.append(["Date & time:", "", "", "Date & time:"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pending_xlsx(orders: Iterable[OmsOrder], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pending"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    cols = [("S.No", 6), ("Sales Channel", 28), ("AWB / Tracking", 22), ("Courier", 16), ("Order ID", 22),
            ("Invoice No", 20), ("SKUs", 30), ("Qty", 6), ("Payment", 10), ("Order Date", 16), ("SLA", 16),
            ("OMS Status", 18)]
    header_row = ws.max_row + 1
    ws.append([c for c, _ in cols])
    for i, o in enumerate(orders, start=1):
        ws.append([
            i, o.channel.name if o.channel else o.channel_label, o.tracking_raw, o.shipping_company,
            o.channel_order_id, o.invoice_id, _skus(o), o.total_qty, o.order_type,
            to_local(o.order_date).strftime("%d-%m-%Y %H:%M") if o.order_date else "",
            to_local(o.sla_date).strftime("%d-%m-%Y %H:%M") if o.sla_date else "", o.status_text,
        ])
    _style_sheet(ws, [w for _, w in cols], header_row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


OMS_DISPATCH_COLUMNS = ["Channel", "Channel Order ID", "Channel Sub Order ID", "Invoice No", "AWB",
                        "Courier", "Dispatch Date", "Dispatch Time"]


def oms_dispatch_csv(scans: Iterable[Scan]) -> str:
    """One line per sub-order, for bulk-marking shipments as dispatched in OMSGuru.

    OMSGuru's public API has no dispatch endpoint, so this file is the bridge. Adjust the column
    names here to match the exact bulk-upload template OMSGuru gives you.
    """
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(OMS_DISPATCH_COLUMNS)
    for s in scans:
        od = _od(s)
        local = to_local(s.scanned_at)
        subs = [x for x in od.get("sub_order_ids") or [] if x] or [""]
        for sub in subs:
            w.writerow([
                s.channel.name if s.channel else "", od.get("channel_order_id", ""), sub,
                od.get("invoice_id", ""), s.tracking_raw, od.get("courier", ""),
                local.strftime("%d-%m-%Y"), local.strftime("%H:%M:%S"),
            ])
    return out.getvalue()


BUCKET_LABELS = {
    "pending": "Pending (due today)", "overdue": "Overdue (AWB from an earlier day)",
    "left_unscanned": "Left Ready-to-ship in OMS without a scan", "cancelled": "Cancelled / returned after AWB",
    "scanned": "Scanned", "generated": "All AWBs generated",
}


def reconcile_xlsx(rows: list[dict], title: str, summary: dict | None = None) -> bytes:
    """Rows from reconcile.rows_payload: AWB-level reconciliation list."""
    wb = Workbook()
    ws = wb.active
    ws.title = "AWBs"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    cols = [("S.No", 6), ("Sales Channel", 28), ("AWB", 22), ("Courier", 16), ("Order ID", 22), ("Invoice No", 20),
            ("SKUs", 30), ("Qty", 6), ("Payment", 10), ("AWB generated", 16), ("Age (days)", 9), ("SLA", 16),
            ("OMS status", 18), ("Reconciliation", 30), ("Scanned at", 18), ("Scanned by", 16)]
    header_row = ws.max_row + 1
    ws.append([c for c, _ in cols])
    late = PatternFill("solid", fgColor="FEE2E2")
    for i, r in enumerate(rows, start=1):
        od = r.get("order") or {}
        sc = r.get("scan") or {}
        ws.append([
            i, (sc.get("channel_name") if sc else "") or od.get("channel_label", ""), r.get("awb", ""),
            od.get("courier", ""), od.get("channel_order_id", ""), od.get("invoice_id", ""), _od_skus(od),
            od.get("total_qty", ""), od.get("order_type", ""), r.get("awb_generated_local", ""), r.get("age_days", ""),
            to_local(datetime.fromisoformat(od["sla_date"].replace("Z", ""))).strftime("%d-%m-%Y %H:%M") if od.get("sla_date") else "",
            od.get("status_text", ""), BUCKET_LABELS.get(r.get("bucket", ""), r.get("bucket", "")),
            sc.get("scanned_at_local", "") if sc else "", sc.get("user", "") if sc else "",
        ])
        if r.get("bucket") == "overdue" or r.get("sla_breached"):
            for c in ws[ws.max_row]:
                c.fill = late
    _style_sheet(ws, [w for _, w in cols], header_row)
    if summary:
        s2 = wb.create_sheet("By channel")
        head = ["Sales Channel", "AWBs generated", "Scanned", "Pending", "Overdue", "Left RTS without scan", "Cancelled", "Scanned %"]
        s2.append(head)
        for ch in summary.get("channels", []):
            s2.append([ch["name"], ch["generated"], ch["scanned"], ch["pending"], ch["overdue"], ch["left_unscanned"],
                       ch["cancelled"], ch["pct"] if ch["pct"] is not None else ""])
        t = summary.get("totals", {})
        s2.append(["TOTAL", t.get("generated"), t.get("scanned"), t.get("pending"), t.get("overdue"),
                   t.get("left_unscanned"), t.get("cancelled"), t.get("pct") if t.get("pct") is not None else ""])
        for c in s2[s2.max_row]:
            c.font = Font(bold=True)
        _style_sheet(s2, [30, 14, 10, 10, 10, 20, 10, 10], 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def channel_summary_xlsx(periods: list[str], channels: list[dict], cells: dict, title: str) -> bytes:
    """Long-term channel-wise scan totals: one row per day / month, one column per sales channel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Scanned by channel"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Period"] + [c["name"] for c in channels] + ["Total"])
    for p in periods:
        vals = [cells.get((p, c["id"]), {}).get("scanned", 0) for c in channels]
        ws.append([p] + vals + [sum(vals)])
    totals = [sum(cells.get((p, c["id"]), {}).get("scanned", 0) for p in periods) for c in channels]
    ws.append(["TOTAL"] + totals + [sum(totals)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    _style_sheet(ws, [14] + [18] * len(channels) + [12], header_row)

    d = wb.create_sheet("Detail")
    cols = ["Period", "Sales Channel", "Scanned", "Verified", "Need check", "Unverified", "Alerts after scan",
            "Duplicate tries", "Wrong marketplace tries", "Blocked (cancelled / return)"]
    d.append(cols)
    for p in periods:
        for c in channels:
            v = cells.get((p, c["id"]))
            if not v:
                continue
            d.append([p, c["name"], v.get("scanned", 0), v.get("OK", 0), v.get("WARN", 0), v.get("UNVERIFIED", 0),
                      v.get("alerts", 0), v.get("DUPLICATE", 0), v.get("WRONG_CHANNEL", 0), v.get("BLOCKED", 0)])
    _style_sheet(d, [12, 28, 10, 10, 11, 11, 14, 14, 18, 20], 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
